#!/usr/bin/env python3
"""eia861_dg_capacity_share.py — FUSION HYPOTHESIS (b)'s gate-1 (DATA)
follow-up to SWPP solar's residual overshoot, a SECOND, independent test of
mechanism (ii) (BTM/distributed solar folded into EIA-930's SUN total)
alongside eia_dpv_btm_solar_share.py's generation-based check.

BACKGROUND (research/open_questions.md, FUSION HYPOTHESES (b)):
eia_dpv_btm_solar_share.py used EIA-923's small-scale-solar GENERATION
series (DPV) and found DPV's share of state solar generation covers
~67% of SWPP solar's observed overshoot (1.1843x implied vs. 1.275x-1.286x
observed) — "not ruled out, not fully sufficient alone." That check is
generation-denominated; the gate-1 test itself is CAPACITY-denominated
(registry nameplate MW vs. EIA-930's max hourly generation MWh). This
script closes that unit mismatch by finding an independent CAPACITY
figure for behind-the-meter/distributed PV, from a source the DPV check
never touched: Form EIA-861's Net Metering and Non-Net-Metering
Distributed Generation tables (a state utility-regulation survey, not a
generation-metering one — genuinely independent of both EIA-930 and
EIA-923).

WHAT THIS SCRIPT DOES: sums EIA-861's own reported Photovoltaic CAPACITY
(net-metered + non-net-metered distributed, both real reporting
categories, summed because a rooftop system can be interconnected either
way and both are physically capacity NOT in the EIA-860 utility-scale
registry) for SWPP's three least-cross-BA-ambiguous core states (OK, KS,
NE — same scope choice eia_dpv_btm_solar_share.py made, for the same
reason: the other 11 SWPP-footprint states straddle multiple BAs and
would reintroduce the attribution ambiguity this thread already spent two
sessions resolving for the registry side). Compares that DG capacity
against EIA-860's own utility-scale solar nameplate capacity for the same
three states, giving a capacity share and, via the same 1/(1-share)
transform the DPV script uses, an implied overshoot ratio directly
comparable (same units) to the observed gate-1 ratio.

HONESTY CAVEATS, stated up front (REASONING STANDARD #4/#7/#10) — this is
weaker evidence than it looks at first glance, for three independent
reasons, none of which the DPV check carries:

1. DATA VINTAGE: Form EIA-861's 2025 data is an EARLY RELEASE (published
   August 2026). Every row in its "States- State Level" sheets carries
   EIA's own caveat verbatim: "The data has not been fully edited and is
   inappropriate for aggregation, such as to state or national totals."
   This script does exactly the aggregation EIA warns against. Read as a
   plausibility check, not a certified figure — final, fully-edited 2025
   EIA-861 data is not out yet ("later in 2026" per the same notice).
2. DOUBLE-COUNTING RISK: EIA-861's Non-Net-Metering table separately
   reports "Capacity Utility Owned" (all-technology, not PV-specific) —
   utility-owned distributed generation that MIGHT already have a row in
   EIA-860's own utility-scale registry (this script's other input),
   which would double-count that slice. This script cannot isolate the
   PV-specific utility-owned figure (the table only breaks "utility
   owned" out by ALL technologies combined, not per-fuel) — it reports
   the all-technology utility-owned share as a bound on how large this
   risk could be, not a precise correction.
3. CAPACITY SHARE != GENERATION SHARE: distributed/rooftop PV typically
   has a lower capacity factor at any given hour than utility-scale
   tracking arrays (fixed tilt, worse orientation diversity, some
   west-facing systems catching the evening rather than the midday peak
   gate-1's MAX-hour metric cares about) — so a capacity share is not
   guaranteed to equal the corresponding generation share at that peak
   hour. This script's implied-overshoot number is therefore a distinct,
   complementary estimate, not a strictly superior replacement for the
   DPV script's generation-based one; both being directionally consistent
   is the useful signal, not either number in isolation.

Files (not fetched automatically — same manual-download precedent every
sibling EIA-860/861 script in this directory documents):
    curl -L -A "Mozilla/5.0" -o /tmp/eia861.zip \
        https://www.eia.gov/electricity/data/eia861/zip/f8612025er.zip
    unzip -o /tmp/eia861.zip -d /tmp/eia861 \
        "Net_Metering_2025_Data_Early_Release.xlsx" \
        "Non_Net_Metering_Distributed_2025_Data_Early_Release.xlsx"
    curl -L -A "Mozilla/5.0" -o /tmp/eia860.zip \
        https://www.eia.gov/electricity/data/eia860/xls/eia8602025.zip
    unzip -o /tmp/eia860.zip -d /tmp/eia860 \
        "2___Plant_Y2025.xlsx" "3_3_Solar_Y2025.xlsx"
    python3 scripts/eia861_dg_capacity_share.py \
        --net-metering /tmp/eia861/Net_Metering_2025_Data_Early_Release.xlsx \
        --non-net-metering /tmp/eia861/Non_Net_Metering_Distributed_2025_Data_Early_Release.xlsx \
        --plants /tmp/eia860/2___Plant_Y2025.xlsx \
        --solar /tmp/eia860/3_3_Solar_Y2025.xlsx \
        --overshoot-ratio 1.286
"""
import argparse
import json
import sys

DEFAULT_STATES = ("OK", "KS", "NE")  # same SWPP core-state scope choice as eia_dpv_btm_solar_share.py

# "States- State Level" sheet, row index 3 (0-based) is the real header;
# column indices below verified live this session against the actual
# 2025 early-release column layout (both sheets carry repeated per-
# technology/per-ownership column blocks that a plain header.index() name
# lookup cannot disambiguate — position confirmed by hand against a known
# state's own totals, not guessed).
NET_METERING_YEAR_COL = 1
NET_METERING_STATE_COL = 2
NET_METERING_PV_TOTAL_CAPACITY_COL = 7  # Photovoltaic / Capacity MW / Total (Res+Comm+Ind+Trans)

NON_NET_METERING_YEAR_COL = 1
NON_NET_METERING_STATE_COL = 2
NON_NET_METERING_ALL_TECH_UTILITY_OWNED_COL = 6  # all-technology, NOT PV-specific — see caveat 2 above
NON_NET_METERING_ALL_TECH_TOTAL_CAPACITY_COL = 4
NON_NET_METERING_PV_TOTAL_CAPACITY_COL = 13  # Photovoltaic (MW) / Total (Res+Comm+Ind+Trans+DirectConnected)


def _numeric(value):
    """EIA's early-release convention: '.' marks a suppressed/not-applicable
    cell, not a true zero, but for a CAPACITY SUM across states this script
    treats it as 0 contribution (consistent with eia_dpv_btm_solar_share.py's
    own None-handling) rather than silently dropping the state — a missing
    technology in one state is not grounds to exclude that state's other
    real capacity."""
    if value is None or value == ".":
        return 0.0
    return float(value)


def find_state_level_rows(rows, year, states, year_col, state_col):
    """rows: raw (values_only) row tuples from a 'States- State Level'
    sheet, including the title/header rows. Returns {state: row} for the
    requested year and states only — every sibling row (other years, other
    states, the caveat/header rows whose year_col isn't an int) is
    filtered out here so callers never have to re-derive the state filter
    themselves."""
    out = {}
    for row in rows:
        if row is None or len(row) <= max(year_col, state_col):
            continue
        if row[year_col] != year:
            continue
        state = row[state_col]
        if state in states:
            out[state] = row
    return out


def net_metering_pv_capacity_mw(row):
    """A 'States- State Level' Net Metering row -> total PV net-metered
    capacity MW (Residential+Commercial+Industrial+Transportation)."""
    return _numeric(row[NET_METERING_PV_TOTAL_CAPACITY_COL])


def non_net_metering_pv_capacity_mw(row):
    """A 'States- State Level' Non-Net-Metering row -> total PV
    non-net-metered distributed capacity MW."""
    return _numeric(row[NON_NET_METERING_PV_TOTAL_CAPACITY_COL])


def combined_dg_pv_capacity(net_metering_rows, non_net_metering_rows, states):
    """Sums net-metered + non-net-metered PV capacity per state (both are
    real, distinct distributed-generation interconnection categories that
    together approximate total behind-the-meter/distributed PV — neither
    alone is the full picture) and in aggregate across `states`."""
    per_state = {}
    total = 0.0
    for state in states:
        nm_row = net_metering_rows.get(state)
        nnm_row = non_net_metering_rows.get(state)
        nm_mw = net_metering_pv_capacity_mw(nm_row) if nm_row is not None else None
        nnm_mw = non_net_metering_pv_capacity_mw(nnm_row) if nnm_row is not None else None
        if nm_mw is None and nnm_mw is None:
            per_state[state] = None
            continue
        state_total = (nm_mw or 0.0) + (nnm_mw or 0.0)
        per_state[state] = {
            "net_metering_mw": round(nm_mw, 3) if nm_mw is not None else None,
            "non_net_metering_mw": round(nnm_mw, 3) if nnm_mw is not None else None,
            "total_mw": round(state_total, 3),
        }
        total += state_total
    return per_state, round(total, 3)


def all_tech_utility_owned_bound(non_net_metering_rows, states):
    """Bounds the double-counting risk (caveat 2): all-technology
    'Capacity Utility Owned' as a fraction of all-technology 'Total
    Capacity' in the Non-Net-Metering table, per state and pooled. This is
    NOT a PV-specific figure (the source table does not break utility-
    ownership out by fuel) — it is reported as the loosest honest bound on
    how large a PV-specific double-count COULD be if utility ownership
    were uniformly distributed across technologies, not a claim that it
    is."""
    total_owned = 0.0
    total_capacity = 0.0
    per_state = {}
    for state in states:
        row = non_net_metering_rows.get(state)
        if row is None:
            per_state[state] = None
            continue
        owned = _numeric(row[NON_NET_METERING_ALL_TECH_UTILITY_OWNED_COL])
        cap = _numeric(row[NON_NET_METERING_ALL_TECH_TOTAL_CAPACITY_COL])
        per_state[state] = {"utility_owned_mw": round(owned, 3), "all_tech_total_mw": round(cap, 3)}
        total_owned += owned
        total_capacity += cap
    pooled_share = round(total_owned / total_capacity, 4) if total_capacity else None
    return per_state, pooled_share


def utility_scale_solar_capacity_by_state(solar_generator_rows, plant_state_by_code):
    """solar_generator_rows: iterable of (status, plant_code, nameplate_mw)
    from EIA-860 Schedule 3.3 (Solar). Sums OPERATING nameplate capacity
    per state, using plant_state_by_code (Schedule 2, Plant file) to map
    plant code -> state. Same OP-status-only convention as
    eia860_missing_plants_check.py's eia860_capacity_by_code."""
    from collections import defaultdict
    out = defaultdict(float)
    unmatched_plants = 0
    for status, code, nameplate_mw in solar_generator_rows:
        if status != "OP":
            continue
        state = plant_state_by_code.get(code)
        if state is None:
            unmatched_plants += 1
            continue
        out[state] += nameplate_mw or 0.0
    return {state: round(mw, 3) for state, mw in out.items()}, unmatched_plants


def dg_capacity_share(dg_mw, utility_scale_mw):
    """DG's share of (DG + utility-scale) capacity. None if both are zero
    (undefined, not zero-share)."""
    denom = dg_mw + utility_scale_mw
    if denom <= 0:
        return None
    return round(dg_mw / denom, 4)


def implied_overshoot_from_capacity_share(share):
    """Same transform as eia_dpv_btm_solar_share.implied_overshoot_from_btm_share:
    if EIA-930's reported generation total folds in DG at capacity share s
    while the registry denominator is utility-scale-only, the implied
    overshoot is 1/(1-s). Reimplemented here (not imported) — this module
    is deliberately self-contained, same convention as every sibling
    EIA-860/861/930 script in this directory."""
    if share is None or share >= 1:
        return None
    return round(1 / (1 - share), 4)


# ── I/O (excluded from unit tests, exercised only by running this file directly) ──

def load_state_level_rows(xlsx_path, sheet):
    import openpyxl  # session-run only, same convention as every sibling script
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet]
    return list(ws.iter_rows(values_only=True))


def load_eia860_plant_directory(xlsx_path):
    """EIA-860 Schedule 2 (Plant file): Plant Code -> State. Same
    header-lookup-by-name convention as eia860_add_missing_plants.py's
    identically-named function."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    next(rows)  # title row
    hdr = next(rows)
    i_code = hdr.index("Plant Code")
    i_state = hdr.index("State")
    out = {}
    for row in rows:
        if row is None or row[i_code] is None:
            continue
        out[row[i_code]] = row[i_state]
    return out


def load_eia860_solar_generator_rows(xlsx_path):
    """Same convention as eia860_missing_plants_check.py's
    load_eia860_generator_rows."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    next(rows)
    hdr = next(rows)
    i_status = hdr.index("Status")
    i_code = hdr.index("Plant Code")
    i_cap = hdr.index("Nameplate Capacity (MW)")
    for row in rows:
        if row is None or row[i_status] is None:
            continue
        yield (row[i_status], int(row[i_code]), row[i_cap])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--net-metering", required=True, help="path to EIA-861 Net_Metering_2025_Data_Early_Release.xlsx")
    ap.add_argument("--non-net-metering", required=True, help="path to EIA-861 Non_Net_Metering_Distributed_2025_Data_Early_Release.xlsx")
    ap.add_argument("--plants", required=True, help="path to EIA-860 2___Plant_Y2025.xlsx")
    ap.add_argument("--solar", required=True, help="path to EIA-860 3_3_Solar_Y2025.xlsx")
    ap.add_argument("--year", type=int, default=2025)
    ap.add_argument("--states", default=",".join(DEFAULT_STATES))
    ap.add_argument("--overshoot-ratio", type=float, default=None,
                     help="the gate-1 overshoot ratio to compare against (e.g. SWPP solar's 1.275-1.286x across sessions)")
    args = ap.parse_args()

    states = tuple(s.strip() for s in args.states.split(",") if s.strip())

    nm_raw = load_state_level_rows(args.net_metering, "States- State Level")
    nnm_raw = load_state_level_rows(args.non_net_metering, "States- State Level")
    nm_rows = find_state_level_rows(nm_raw, args.year, states, NET_METERING_YEAR_COL, NET_METERING_STATE_COL)
    nnm_rows = find_state_level_rows(nnm_raw, args.year, states, NON_NET_METERING_YEAR_COL, NON_NET_METERING_STATE_COL)
    missing_states = [s for s in states if s not in nm_rows or s not in nnm_rows]
    if missing_states:
        print(f"[eia861_dg_capacity_share] no {args.year} row for: {missing_states}", file=sys.stderr)

    dg_per_state, dg_total_mw = combined_dg_pv_capacity(nm_rows, nnm_rows, states)
    owned_bound_per_state, owned_bound_pooled_share = all_tech_utility_owned_bound(nnm_rows, states)

    plant_state = load_eia860_plant_directory(args.plants)
    util_by_state, unmatched = utility_scale_solar_capacity_by_state(
        load_eia860_solar_generator_rows(args.solar), plant_state)
    if unmatched:
        print(f"[eia861_dg_capacity_share] {unmatched} EIA-860 solar generator row(s) had no plant-directory match", file=sys.stderr)
    util_total_mw = round(sum(util_by_state.get(s, 0.0) for s in states), 3)

    share = dg_capacity_share(dg_total_mw, util_total_mw)
    implied = implied_overshoot_from_capacity_share(share)

    report = {
        "check": "eia861_dg_capacity_share",
        "note": ("second, capacity-denominated, independent magnitude check for FUSION "
                 "HYPOTHESIS (b) mechanism (ii) — complementary to eia_dpv_btm_solar_share.py's "
                 "generation-based check, NOT a replacement or a confirmation. See module "
                 "docstring for the EIA-861-early-release / double-counting / capacity-vs-"
                 "generation-share caveats before trusting this number."),
        "year": args.year,
        "states": list(states),
        "dg_pv_capacity_mw": {"per_state": dg_per_state, "total_mw": dg_total_mw},
        "utility_scale_solar_capacity_mw": {"per_state": {s: util_by_state.get(s) for s in states}, "total_mw": util_total_mw},
        "all_tech_utility_owned_bound": {"per_state": owned_bound_per_state, "pooled_share_of_non_net_metering_capacity": owned_bound_pooled_share},
        "dg_capacity_share": share,
        "implied_overshoot_ratio_if_sun_includes_dg": implied,
        "observed_gate1_overshoot_ratio": args.overshoot_ratio,
    }
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
